from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from draw_budget_state_flow import draw_budget_state_flow


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "文档/迭代/预算优化/详细需求说明.md"
OUTPUT = ROOT / "文档/迭代/预算优化/预算优化详细需求说明_V1.11.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, name="Aptos", size=10.5, bold=None, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def style_document(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(1.8)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)
    sec.header_distance = Cm(0.8)
    sec.footer_distance = Cm(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.22

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(12)
    title_ppr = title._element.get_or_add_pPr()
    title_borders = title_ppr.find(qn("w:pBdr"))
    if title_borders is not None:
        title_ppr.remove(title_borders)

    for name, size, space_before, space_after in [
        ("Heading 1", 16, 16, 7),
        ("Heading 2", 13, 12, 5),
        ("Heading 3", 11.5, 9, 4),
    ]:
        st = styles[name]
        st.font.name = "Aptos Display"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.space_before = Pt(space_before)
        st.paragraph_format.space_after = Pt(space_after)
        st.paragraph_format.keep_with_next = True

    toc_style = styles['TOC 1'] if 'TOC 1' in styles else styles.add_style('TOC 1', WD_STYLE_TYPE.PARAGRAPH)
    toc_style.font.name = 'Aptos'
    toc_style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    toc_style.font.size = Pt(10.5)
    toc_style.font.color.rgb = RGBColor(0, 0, 0)
    toc_style.paragraph_format.space_before = Pt(0)
    toc_style.paragraph_format.space_after = Pt(3)
    toc_style.paragraph_format.line_spacing = 1.1
    toc_style.paragraph_format.tab_stops.add_tab_stop(Cm(16.6), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("ERP-I003 预算优化 | 研发与测试需求基线")
    set_font(run, size=8.5)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("ERP-I003 预算优化  ·  ")
    set_font(run, size=8.5, color="666666")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)


def add_inline(paragraph, text, inherit=False):
    text = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", text)
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, name="Consolas", size=9.5)
        else:
            run = paragraph.add_run(part)
            if not inherit:
                set_font(run)


def add_body(doc, text, style=None, indent=0):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.left_indent = Cm(indent) if indent else None
    p.paragraph_format.line_spacing = 1.18
    if text.startswith(("**已确认的转换口径**", "“正式提交成功”指")):
        p.paragraph_format.keep_together = True
    add_inline(p, text)
    return p


def add_bullet(doc, text, level=0, ordered=False):
    style = "List Number" if ordered else "List Bullet"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.left_indent = Cm(0.55 + level * 0.45)
    p.paragraph_format.first_line_indent = Cm(-0.25)
    p.paragraph_format.space_after = Pt(2)
    add_inline(p, text)
    return p


def add_table(doc, rows):
    if not rows:
        return
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    if len(rows[0]) == 2:
        widths = [3.5, 13.1]
    elif rows[0][0] == "场景步骤":
        widths = [8.0, 2.8, 2.8, 3.0]
    elif rows[0][0] == "单据状态" and len(rows[0]) == 4:
        widths = [2.1, 6.1, 4.2, 4.2]
    elif len(rows[0]) == 4:
        widths = [1.3, 7.0, 3.3, 5.0]
    elif rows[0][0] == "预算状态":
        widths = [2.5, 6.0, 8.1]
    elif rows[0][0] == "截止时点的有效情况":
        widths = [7.0, 4.8, 4.8]
    elif rows[0][0] == "阶段":
        widths = [2.5, 5.4, 8.7]
    elif rows[0][0] == "单据类型" and len(rows[0]) == 3:
        widths = [4.0, 5.0, 7.6]
    elif rows[0][0] == "节点或事件":
        widths = [3.4, 9.9, 3.3]
    elif rows[0][0] == "单据类型":
        widths = [3.1, 3.4, 10.1]
    elif rows[0][0] == "ID":
        widths = [1.4, 3.8, 11.4]
    elif rows[0][0] == "版本":
        widths = [1.7, 2.7, 12.2]
    else:
        widths = [4.2, 6.2, 6.2]
    for col, width in zip(table.columns, widths):
        col.width = Cm(width)
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, value in enumerate(rows[0]):
        cell = header.cells[i]
        cell.width = Cm(widths[i])
        cell.text = ""
        set_cell_shading(cell, "244054")
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.keep_with_next = True
        run = p.add_run(value.strip())
        set_font(run, size=9.2, bold=True, color="FFFFFF")
    for row_index, row_data in enumerate(rows[1:], start=1):
        cells = table.add_row().cells
        tr_pr = table.rows[-1]._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        for i, value in enumerate(row_data):
            cell = cells[i]
            cell.width = Cm(widths[i])
            cell.text = ""
            set_cell_shading(cell, "F5F7FA" if row_index % 2 == 0 else "FFFFFF")
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.08
            if rows[0][0] == "预算状态":
                p.paragraph_format.keep_with_next = row_index < len(rows) - 1
                p.paragraph_format.keep_together = True
            add_inline(p, value.strip())
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_figure(doc, label, relative_path):
    image_path = (SOURCE.parent / relative_path).resolve()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    picture = p.add_run().add_picture(str(image_path), width=Cm(16.6))
    picture._inline.docPr.set("descr", label)
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(7)
    set_font(caption.add_run(label), size=9, color="4B5563")


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        raw = lines[i].strip().strip("|")
        cells = [c.strip() for c in raw.split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells):
            rows.append(cells)
        i += 1
    return rows, i


def add_contents(doc):
    heading = doc.add_paragraph()
    heading.paragraph_format.space_before = Pt(12)
    heading.paragraph_format.space_after = Pt(7)
    heading.paragraph_format.keep_with_next = True
    set_font(heading.add_run('目录'), size=16, bold=True)
    p = doc.add_paragraph()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instruction = OxmlElement('w:instrText')
    instruction.set(qn('xml:space'), 'preserve')
    instruction.text = ' TOC \\o "1-1" \\h \\z '
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    p.add_run()._r.append(begin)
    p.add_run()._r.append(instruction)
    p.add_run()._r.append(separate)
    p.add_run('更新目录以显示章节和页码。')
    p.add_run()._r.append(end)


def build():
    source_text = SOURCE.read_text(encoding="utf-8")
    source_lines = source_text.splitlines()
    version = re.search(r"^版本：(\S+)", source_text, re.MULTILINE).group(1)
    date = re.search(r"^整理日期：(\d{4})-(\d{2})-(\d{2})", source_text, re.MULTILINE)
    date_label = f"{date.group(1)}年{int(date.group(2))}月{int(date.group(3))}日"
    draw_budget_state_flow()
    doc = Document()
    style_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_inline(title, "ERP I003 预算优化详细需求说明", inherit=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(15)
    run = subtitle.add_run(f"研发与测试评审版  |  {version}  |  {date_label}")
    set_font(run, size=11, color="4B5563")

    add_contents(doc)

    i = next(i for i, line in enumerate(source_lines) if line.startswith("## "))
    code_block = False
    code_lines = []
    code_type = ""
    while i < len(source_lines):
        line = source_lines[i].rstrip()
        if line.startswith("```"):
            if not code_block:
                code_block = True
                code_lines = []
                code_type = line[3:].strip()
            else:
                if code_type:
                    p = doc.add_paragraph()
                    p.paragraph_format.left_indent = Cm(0.4)
                    p.paragraph_format.right_indent = Cm(0.4)
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(6)
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.keep_together = True
                    run = p.add_run("\n".join(code_lines))
                    set_font(run, name="Consolas", size=9, color="374151")
                code_block = False
            i += 1
            continue
        if code_block:
            code_lines.append(line)
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        image_match = re.fullmatch(r"!\[([^\]]+)\]\(([^)]+)\)", line)
        if image_match:
            add_figure(doc, image_match.group(1), image_match.group(2))
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(source_lines) and source_lines[i + 1].strip().startswith("|"):
            rows, i = parse_table(source_lines, i)
            add_table(doc, rows)
            continue
        m = re.match(r"^(#{2,4})\s+(.*)$", line)
        if m:
            level = len(m.group(1)) - 1
            text = re.sub(r"^\d+(?:\.\d+)*\s+", lambda x: x.group(0), m.group(2))
            p = doc.add_paragraph(style=f"Heading {level}")
            if m.group(2).startswith(("2.2 ", "6. ")):
                p.paragraph_format.page_break_before = True
            add_inline(p, text, inherit=True)
            i += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            p = add_body(doc, line)
            p.paragraph_format.left_indent = Cm(0.45)
            p.paragraph_format.first_line_indent = Cm(-0.45)
            i += 1
            continue
        if line.startswith("- "):
            add_bullet(doc, line[2:], ordered=False)
            i += 1
            continue
        add_body(doc, line)
        i += 1

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "ERP-I003 预算优化详细需求说明"
    doc.core_properties.subject = f"预算优化研发与测试需求基线 {version}"
    doc.core_properties.author = "ERP 项目组"
    doc.core_properties.keywords = "ERP-I003,预算优化,需求,测试"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
