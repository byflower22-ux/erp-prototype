from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

SKILL_DIR = Path(r"C:\Users\BS\.codex\plugins\cache\openai-primary-runtime\documents\26.819.11345\skills\documents")
sys.path.insert(0, str(SKILL_DIR / "scripts"))
from table_geometry import apply_table_geometry

ROOT = Path(__file__).parent
SOURCE = ROOT / "付款报销回单对接用户操作手册.md"
OUTPUT = ROOT / "付款报销回单对接用户操作手册.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2937"
MUTED = "5B6573"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"


def set_font(run, size=11, bold=None, color=INK):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, *, bold=False, color=INK, size=9.2, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.1
    r = p.add_run(str(text))
    set_font(r, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        shade(table.rows[0].cells[i], LIGHT_BLUE)
        set_cell_text(table.rows[0].cells[i], header, bold=True, color=DARK_BLUE, size=9.2,
                      align=WD_ALIGN_PARAGRAPH.CENTER)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=9.0)
    # Content-aware fixed widths, always summing to 9360 DXA.
    weights = [max(10, len(str(h)) * 1.2 + max((len(str(r[i])) for r in rows), default=0) * 0.32)
               for i, h in enumerate(headers)]
    total = sum(weights)
    widths = [max(1150, int(9360 * w / total)) for w in weights]
    delta = 9360 - sum(widths)
    widths[-1] += delta
    apply_table_geometry(table, widths, table_width_dxa=9360, indent_dxa=120)
    header_pr = table.rows[0]._tr.get_or_add_trPr()
    header_flag = OxmlElement("w:tblHeader")
    header_flag.set(qn("w:val"), "true")
    header_pr.append(header_flag)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def configure(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        st = doc.styles[name]
        st.font.name = "Calibri"
        st._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("付款报销回单对接 | 财务人员操作手册")
    set_font(r, size=8.5, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("内部使用")
    set_font(r, size=8.5, color=MUTED)


def add_rich_text(doc, text, style="Normal", callout=False):
    p = doc.add_paragraph(style=style)
    if callout:
        p.paragraph_format.left_indent = Inches(0.12)
        p.paragraph_format.right_indent = Inches(0.12)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(7)
        ppr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), CALLOUT)
        ppr.append(shd)
    # Bold markdown labels and inline code are rendered as readable text.
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            set_font(r, size=10.5, bold=True, color=DARK_BLUE)
        else:
            value = part[1:-1] if part.startswith("`") and part.endswith("`") else part
            r = p.add_run(value)
            set_font(r, size=10.5, color=INK)
    return p


def build():
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure(doc)
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(line[2:].strip())
            set_font(r, size=24, bold=True, color="0B2545")
            i += 1
            continue
        if line.startswith("## ") or line.startswith("### "):
            level = 1 if line.startswith("## ") else 2
            add_rich_text(doc, line[(3 if level == 1 else 4):].strip(), style=f"Heading {level}")
            i += 1
            continue
        if line.startswith("> "):
            add_rich_text(doc, line[2:].strip(), callout=True)
            i += 1
            continue
        if line.startswith("| "):
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                table_lines.append(lines[i])
                i += 1
            if len(table_lines) >= 2:
                parse = lambda s: [x.strip() for x in s.strip().strip("|").split("|")]
                headers = parse(table_lines[0])
                rows = [parse(x) for x in table_lines[2:]]
                add_table(doc, headers, rows)
            continue
        if re.match(r"^\d+\. ", line):
            add_rich_text(doc, re.sub(r"^\d+\. ", "", line), style="List Number")
            i += 1
            continue
        if line.startswith("- [ ] "):
            add_rich_text(doc, "[ ] " + line[6:].strip())
            i += 1
            continue
        if line.startswith("- "):
            add_rich_text(doc, line[2:].strip(), style="List Bullet")
            i += 1
            continue
        add_rich_text(doc, line.strip())
        i += 1
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
